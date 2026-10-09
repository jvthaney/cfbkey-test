#!/usr/bin/env python3
"""
draft_alt_text.py

Drafts alt text for the pictures and detail figures in the workbook, so reviewers
start from something instead of a blank cell.

  python draft_alt_text.py --workbook "$ZK/cfbkey_workbook.xlsx"

WHAT IT DOES
  * Adds two columns at the END of the 'images' and 'figures' sheets:
        alt_text_draft   the suggested wording
        draft_basis      how it was made (so you can review the weakest first)
  * Never touches 'alt_text_new' (the column for reviewed text), so your edits are safe.
  * Makes a dated backup copy of the workbook first.
  * Leaves existing drafts alone unless you add --force.
  * Adds two progress lines to the README sheet.

HOW THE DRAFTS ARE MADE (these are guesses, not descriptions of what each picture shows)
  detail figures:  the sentence on the species page that links to the figure
                   ("Antennal setae formula 0-0-1-3 / 1-1-3"), else the link text,
                   else the file name.
  photographs:     the species name plus whatever the file name says
                   (S_serricaudatus_ventral_175x175.jpg -> "..., ventral view"),
                   else a plain "Photograph of <species>".

On the site, reviewed text (alt_text_new) is used first, then the draft, then a plain
fallback. The build report counts drafts separately, so unreviewed text is never
mistaken for finished text.
"""

import argparse
import datetime
import re
import shutil
import sys
from pathlib import Path, PurePosixPath

try:
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("Needs openpyxl:  conda install -c conda-forge openpyxl")

# abbreviations seen in the file names; anything not listed is used as written
ABBR = {
    "pa": "postabdomen", "pc": "postabdominal claw",
    "1st": "first leg", "2nd": "second leg", "3rd": "third leg", "4th": "fourth leg", "5th": "fifth leg",
    "ant": "antenna", "antennae": "antennae", "antennule": "antennule", "antennules": "antennules",
    "valve": "valve", "valves": "valves", "head": "head", "claw": "claw", "legs": "legs", "leg": "leg",
    "carapace": "carapace", "rostrum": "rostrum", "mandible": "mandible", "trophi": "trophi",
    "exopod": "exopod", "endopod": "endopod", "furca": "furca", "caudal": "caudal",
    "gut": "gut", "eye": "eye", "ocellus": "ocellus", "ephippium": "ephippium",
}
VIEWS = {"ventral", "dorsal", "lateral", "anterior", "posterior", "frontal", "side", "top"}
STATES = {"male", "female", "juvenile", "adult", "ephippial", "gravid", "mating"}
SKIP = {"small", "large", "lg", "org", "orig", "roll", "copy", "thumb", "img", "image", "pic", "photo",
        "fig", "figure", "hq", "crop", "cr", "edit", "final", "new", "old", "png", "jpg"}
TOKEN_RE = re.compile(r"\d+x\d+|\d+(?:st|nd|rd|th)(?![a-z])|[A-Za-z]+|\d+")


def plain(md: str) -> str:
    s = re.sub(r"\[([^\]]+)\]\([^)]*\)", r"\1", str(md))
    s = re.sub(r"</?(sup|sub)>", "", s)
    s = s.replace("*", "")
    return re.sub(r"\s+", " ", s).strip()


def cut(text: str, limit: int = 140) -> str:
    text = text.strip().rstrip(".")
    if len(text) <= limit:
        return text
    return text[:limit].rsplit(" ", 1)[0].rstrip(",;:")


def feature_words(path: str, name: str):
    """Meaningful words from a file name, after dropping sizes, numbers and the species name."""
    words = [w.lower() for w in name.replace("(", " ").replace(")", " ").split() if w.strip()]
    genus = words[0] if words else ""
    epithets = words[1:]
    out = []
    for tok in TOKEN_RE.findall(PurePosixPath(path).stem):
        t = tok.lower()
        if re.fullmatch(r"\d+x\d+|\d+", t) or t in SKIP:
            continue
        stripped = re.sub(r"\d+", "", t)
        if (stripped in {genus, genus[:1], "".join(words[:2]), genus[:1] + "".join(epithets)}
                or any(e and (stripped == e or e in stripped) for e in epithets)
                or (genus and len(genus) > 3 and genus in stripped)):
            continue
        t = re.sub(r"^(left|right)(?=[a-z]{4,})", r"\1 ", t)       # leftendopod -> left endopod
        out.append(ABBR.get(t, t))
    return out


def style_header(ws, col, text, width):
    c = ws.cell(row=1, column=col, value=text)
    c.font = Font(name="Arial", bold=True, size=10, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor="7F7F7F")          # grey = generated, not an input
    c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.column_dimensions[get_column_letter(col)].width = width


def ensure_columns(ws):
    heads = [c.value for c in ws[1]]
    for name, width in (("alt_text_draft", 55), ("draft_basis", 28)):
        if name not in heads:
            style_header(ws, ws.max_column + 1, name, width)
            heads = [c.value for c in ws[1]]
    return {h: i + 1 for i, h in enumerate(heads) if h}


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--workbook", required=True)
    ap.add_argument("--force", action="store_true", help="replace existing drafts")
    a = ap.parse_args()
    wbp = Path(a.workbook).expanduser()
    if not wbp.exists():
        sys.exit(f"ERROR: workbook not found: {wbp}")
    stamp = datetime.datetime.now().strftime("%Y%m%d-%H%M%S")
    backup = wbp.with_name(f"{wbp.stem}_backup_{stamp}{wbp.suffix}")
    shutil.copy2(wbp, backup)
    print(f"Backup saved: {backup.name}")

    wb = load_workbook(wbp)
    for need in ("species", "species_sections", "images", "figures"):
        if need not in wb.sheetnames:
            sys.exit(f"ERROR: the workbook has no '{need}' sheet.")

    def rows(name):
        ws = wb[name]
        heads = [c.value for c in ws[1]]
        for r in ws.iter_rows(min_row=2):
            yield r, {h: c.value for h, c in zip(heads, r)}

    names = {d["species_id"]: str(d["display_name"]) for _, d in rows("species") if d.get("species_id")}

    # sentence that links to each detail figure:  (species_id, path) -> (plain text, how many figures it names)
    ctx = {}
    for _, d in rows("species_sections"):
        text = str(d.get("text") or "")
        paths = re.findall(r"\]\(fig:([^)\s]+)\)", text)
        for path in paths:
            ctx.setdefault((d["species_id"], path), (plain(text), len(set(paths))))
    # the link texts used for each figure (a figure can be linked from several places)
    link_texts = {}
    for _, d in rows("figures"):
        t = plain(d.get("caption_on_old_site") or "")
        if t and d.get("species_id") and d.get("image"):
            lst = link_texts.setdefault((d["species_id"], d["image"]), [])
            if t.lower() not in [x.lower() for x in lst]:
                lst.append(t)

    # ---- images ----
    ws = wb["images"]
    col = ensure_columns(ws)
    seen, n_img = {}, {"filename": 0, "generic": 0}
    for r, d in rows("images"):
        sid = d.get("species_id")
        if not sid or not d.get("file"):
            continue
        cell = r[col["alt_text_draft"] - 1]
        if cell.value and not a.force:
            continue
        name = names.get(sid, sid)
        seen[sid] = seen.get(sid, 0) + 1
        feats = feature_words(str(d["file"]), name)
        if feats and all(f in VIEWS for f in feats):
            alt, basis = f"{name}, {' '.join(feats)} view", "filename"
        elif feats and all(f in STATES for f in feats):
            alt, basis = f"{name}, {' '.join(feats)}", "filename"
        elif feats:
            alt, basis = f"{name}: {' '.join(feats)}", "filename"
        elif seen[sid] == 1 and not d.get("on_extra_page"):
            alt, basis = f"Photograph of {name}", "generic"
        else:
            alt, basis = f"Photograph of {name}, view {seen[sid]}", "generic"
        cell.value = alt
        r[col["draft_basis"] - 1].value = basis
        for k in ("alt_text_draft", "draft_basis"):
            r[col[k] - 1].alignment = Alignment(vertical="top", wrap_text=True)
            r[col[k] - 1].font = Font(name="Arial", size=10)
        n_img[basis] += 1
    ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"

    # ---- detail figures ----
    ws = wb["figures"]
    col = ensure_columns(ws)
    n_fig = {"sentence that links to it": 0, "link text": 0,
             "link text (picture linked from several places)": 0, "filename": 0, "generic": 0}
    samples = []
    for r, d in rows("figures"):
        sid, path = d.get("species_id"), d.get("image")
        if not sid or not path:
            continue
        cell = r[col["alt_text_draft"] - 1]
        if cell.value and not a.force:
            continue
        name = names.get(sid, sid)
        sentence, n_named = ctx.get((sid, path), (None, 0))
        caption = plain(d.get("caption_on_old_site") or "")
        shared = link_texts.get((sid, path), [])
        feats = feature_words(str(path), name)
        if len(shared) > 1:
            alt, basis = f"{name}: {' and '.join(shared)}", "link text (picture linked from several places)"
        elif sentence and n_named == 1:
            alt, basis = f"{name}: {cut(sentence)}", "sentence that links to it"
        elif caption:
            alt, basis = f"{name}: {caption}", "link text"
        elif feats:
            alt, basis = f"{name}: {' '.join(feats)}", "filename"
        else:
            alt, basis = f"{name}: detail figure", "generic"
        cell.value = alt
        r[col["draft_basis"] - 1].value = basis
        for k in ("alt_text_draft", "draft_basis"):
            r[col[k] - 1].alignment = Alignment(vertical="top", wrap_text=True)
            r[col[k] - 1].font = Font(name="Arial", size=10)
        n_fig[basis] += 1
        if len(samples) < 6:
            samples.append(alt)
    ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"

    # ---- progress lines on the README ----
    if "README" in wb.sheetnames:
        rd = wb["README"]
        if not any(c.value and "DRAFT alt text" in str(c.value) for row in rd.iter_rows() for c in row[:1]):
            def L(sheet, header):
                heads = [c.value for c in wb[sheet][1]]
                return get_column_letter(heads.index(header) + 1)
            row = rd.max_row + 2
            for label, sheet in (("Images with only a DRAFT alt text (needs review)", "images"),
                                 ("Detail figures with only a DRAFT alt text (needs review)", "figures")):
                d_, n_ = L(sheet, "alt_text_draft"), L(sheet, "alt_text_new")
                rd.cell(row=row, column=1, value=label).font = Font(name="Arial", size=10)
                rd.cell(row=row, column=2,
                        value=f'=COUNTIFS({sheet}!{d_}2:{d_}20000,"<>",{sheet}!{n_}2:{n_}20000,"")'
                        ).font = Font(name="Arial", size=10)
                row += 1

    wb.save(wbp)
    print(f"images:  {sum(n_img.values())} drafted  {n_img}")
    print(f"figures: {sum(n_fig.values())} drafted  {n_fig}")
    for s in samples:
        print("   e.g.", s)
    print("\nThese are starting points. Review them in the workbook (filter 'draft_basis' to see the weakest first:")
    print("'generic' means the file name gave no clue), then write the final wording in alt_text_new.")


if __name__ == "__main__":
    main()
